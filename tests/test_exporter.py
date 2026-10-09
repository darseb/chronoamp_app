"""Unit tests for data.exporter (Excel export)."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
import openpyxl
import pytest

from core.models import DataPoint, MeasurementConfig
from data.exporter import export_to_excel
from data.session_store import save_session
from interpretation.verdict import Verdict


def test_export_to_excel_sheets_and_content(tmp_path: Path) -> None:
    config = MeasurementConfig(potential=0.1, run_time=5.0, interval_time=0.5)
    points = [DataPoint(time=i * 0.5, current=2.0 + i * 0.1) for i in range(10)]
    csv_path = save_session(
        config=config,
        data_points=points,
        verdict=Verdict.POSITIVE,
        explanation="Signal above LOD (2.5000 µA). Predicted concentration: 1.234 nM.",
        sessions_folder=tmp_path / "sessions",
        timestamp=datetime(2026, 9, 13, 12, 0, 0),
    )

    xlsx_path = tmp_path / "exported.xlsx"
    out = export_to_excel(csv_path, xlsx_path)
    assert out.exists()

    wb = openpyxl.load_workbook(out)
    assert "Sheet1" in wb.sheetnames
    assert "Report" in wb.sheetnames

    # Check Sheet1 headers (PSTrace format)
    ws1 = wb["Sheet1"]
    assert ws1.cell(row=1, column=1).value == "CA i vs t"
    assert ws1.cell(row=2, column=1).value == "s"
    assert ws1.cell(row=2, column=2).value == "µA"

    # Check Report sheet
    ws_rep = wb["Report"]
    report_rows = {ws_rep.cell(row=r, column=1).value: ws_rep.cell(row=r, column=2).value for r in range(2, 7)}
    assert report_rows["Verdict"] == "POSITIVE"
    assert "1.234 nM" in report_rows["Explanation"]
